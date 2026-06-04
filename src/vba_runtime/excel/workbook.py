"""VBA Workbook wrapper around openpyxl."""

from __future__ import annotations

from typing import Optional, Union

try:
    import openpyxl
    from openpyxl import Workbook as _OpenpyxlWorkbook
except ImportError:
    openpyxl = None  # type: ignore[assignment]
    _OpenpyxlWorkbook = None  # type: ignore[assignment, misc]

from vba_runtime.excel.worksheet import VBAWorksheet


class VBAWorkbook:
    """Thin wrapper around openpyxl.Workbook providing VBA-like interface.

    Usage:
        wb = VBAWorkbook("report.xlsx")
        ws = wb.sheets("Sheet1")
        val = ws.cells(1, 1).value
        wb.save("output.xlsx")
        wb.close()
    """

    def __init__(self, path: Optional[str] = None) -> None:
        if openpyxl is None:
            raise ImportError(
                "openpyxl is required for Excel support. "
                "Install it with: pip install openpyxl"
            )
        if path:
            self._wb = openpyxl.load_workbook(path)
        else:
            self._wb = openpyxl.Workbook()

    @property
    def active_sheet(self) -> VBAWorksheet:
        """Returns the active worksheet."""
        return VBAWorksheet(self._wb.active)

    def sheets(self, name_or_index: Union[str, int]) -> VBAWorksheet:
        """Access worksheet by name or 1-based index.

        Args:
            name_or_index: Sheet name (str) or 1-based index (int).
        """
        if isinstance(name_or_index, int):
            idx = name_or_index - 1  # Convert to 0-based
            if idx < 0 or idx >= len(self._wb.worksheets):
                raise IndexError(
                    f"Subscript out of range: sheet index {name_or_index}"
                )
            return VBAWorksheet(self._wb.worksheets[idx])
        return VBAWorksheet(self._wb[name_or_index])

    @property
    def sheet_count(self) -> int:
        """Number of worksheets."""
        return len(self._wb.worksheets)

    @property
    def sheet_names(self) -> list[str]:
        """List of worksheet names."""
        return self._wb.sheetnames

    def save(self, path: str) -> None:
        """Save workbook to file."""
        self._wb.save(path)

    def close(self) -> None:
        """Close the workbook."""
        self._wb.close()
